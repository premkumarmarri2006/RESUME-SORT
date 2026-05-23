import io

def extract_text_from_docx(file_bytes: bytes) -> str:
    try:
        from docx import Document
        doc = Document(io.BytesIO(file_bytes))
        lines = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                lines.extend(c.text for c in row.cells if c.text.strip())
        return "\n".join(lines)
    except Exception as e:
        return f"Error: {e}"
